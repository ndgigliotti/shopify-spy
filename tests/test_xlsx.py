"""XLSX output: the CSV's rows as typed cells (issue #35)."""

import json
import sys
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal

import openpyxl
import pytest
import xlsxwriter

import shopify_spy._xlsx as xlsx_module
from shopify_spy._export import ItemWriter, format_for, write_items
from shopify_spy._flatten import COLUMNS, DATE_COLUMNS, ID_COLUMNS, NUMBER_COLUMNS
from shopify_spy._summary import summarize

from .test_export import SHOPIFY_PRODUCT, WOO_PRODUCT


def read_sheets(path):
    """{sheet name: list of row dicts}, plus each sheet's openpyxl worksheet."""
    book = openpyxl.load_workbook(path)
    sheets = {}
    for sheet in book.worksheets:
        header, *rows = sheet.iter_rows(values_only=True)
        sheets[sheet.title] = [dict(zip(header, row, strict=True)) for row in rows]
    return sheets, book


def read_rows(path):
    sheets, _ = read_sheets(path)
    (rows,) = sheets.values()
    return rows


def test_format_detection():
    assert format_for("a.XLSX") == "xlsx"
    assert format_for("a.txt", "xlsx") == "xlsx"


def test_shopify_rows_have_typed_cells(tmp_path):
    path = tmp_path / "out.xlsx"
    product = SHOPIFY_PRODUCT | {
        "product": SHOPIFY_PRODUCT["product"]
        | {"created_at": "2024-01-15T10:30:00-05:00", "published_at": "not a date"}
    }
    assert write_items([product], path) == 1
    first, second = read_rows(path)
    # The CSV's columns, in its order
    assert list(first) == [c for c in COLUMNS if c != "image_paths"]
    # IDs stay text; prices and counts are numbers; booleans are booleans
    assert (first["id"], first["variant_id"]) == ("7", "71")
    assert (first["price"], first["compare_at_price"]) == (10, None)
    assert (second["price"], second["compare_at_price"]) == (12.5, 15)
    assert (first["image_count"], first["variant_count"]) == (2, 2)
    assert (first["available"], second["available"]) == (True, False)
    # Timestamps are dates in UTC; unparseable ones stay text
    assert first["created_at"] == datetime(2024, 1, 15, 15, 30)
    assert first["published_at"] == "not a date"
    assert first["description"] == "Warm & soft Wool"


def test_woocommerce_prices_and_ratings_are_numbers(tmp_path):
    path = tmp_path / "out.xlsx"
    write_items([WOO_PRODUCT | {"average_rating": "4.50", "review_count": 3}], path)
    row, _ = read_rows(path)
    assert (row["price"], row["compare_at_price"], row["min_price"], row["max_price"]) == (
        19.99,
        24.99,
        19.99,
        29.99,
    )
    assert (row["average_rating"], row["review_count"]) == (4.5, 3)
    assert (row["id"], row["variant_id"], row["sku"]) == ("5", "51", "P1")
    assert (row["platform"], row["title"]) == ("woocommerce", "17″ Pan")


def test_both_platforms_share_one_sheet(tmp_path):
    path = tmp_path / "out.xlsx"
    write_items([SHOPIFY_PRODUCT, WOO_PRODUCT], path)
    rows = read_rows(path)
    assert [(r["platform"], r["variant_id"], r["price"]) for r in rows] == [
        ("shopify", "71", 10),
        ("shopify", "72", 12.5),
        ("woocommerce", "51", 19.99),
        ("woocommerce", "52", 19.99),
    ]
    # A column only one platform has is blank for the other
    assert (rows[0]["currency"], rows[2]["currency"]) == (None, "USD")


def test_native_schema_is_refused(tmp_path):
    with pytest.raises(ValueError, match="XLSX is always unified"):
        write_items([SHOPIFY_PRODUCT], tmp_path / "out.xlsx", schema="native")


def test_text_is_never_a_formula(tmp_path):
    path = tmp_path / "out.xlsx"
    titles = ["=1+1", "+cmd", "-2", "@SUM(A1)", '=HYPERLINK("http://x")']
    items = [{"product": {"id": i, "title": t}, "store": "a.com"} for i, t in enumerate(titles)]
    write_items(items, path)
    _, book = read_sheets(path)
    sheet = book.active
    column = [c.value for c in sheet[1]].index("title") + 1
    cells = [sheet.cell(row=r, column=column) for r in range(2, len(titles) + 2)]
    assert [c.value for c in cells] == titles
    assert {c.data_type for c in cells} == {"s"}


def test_numeric_looking_text_stays_text(tmp_path):
    path = tmp_path / "out.xlsx"
    item = {
        "product": {"id": 1, "variants": [{"id": 2, "sku": "00123", "option1": "1e5"}]},
        "store": "a.com",
    }
    write_items([item], path)
    (row,) = read_rows(path)
    assert (row["sku"], row["option1"]) == ("00123", "1e5")


def test_long_text_is_cut_at_excels_limit_in_utf16_units(tmp_path, caplog):
    path = tmp_path / "out.xlsx"
    # Each emoji is two UTF-16 units; an odd limit would split the last one
    description = "😀" * 20_000
    item = {"product": {"id": 1, "body_html": description, "title": "x" * 40_000}}
    write_items([item | {"store": "a.com"}], path)
    (row,) = read_rows(path)
    assert row["description"] == "😀" * (xlsx_module.MAX_TEXT // 2)
    assert row["title"] == "x" * xlsx_module.MAX_TEXT
    assert "description" in caplog.text and "title" in caplog.text


def test_header_is_frozen_and_filterable(tmp_path):
    path = tmp_path / "out.xlsx"
    write_items([SHOPIFY_PRODUCT], path)
    _, book = read_sheets(path)
    sheet = book.active
    assert sheet.title == "items"
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == f"A1:{sheet.cell(row=3, column=sheet.max_column).coordinate}"


def test_full_sheet_continues_on_a_new_one(tmp_path, monkeypatch):
    monkeypatch.setattr(xlsx_module, "MAX_ROWS", 4)  # header + 3 rows
    path = tmp_path / "out.xlsx"
    single = {"product": {"id": 1}, "store": "a.com"}
    # 1 + 2 rows fill the first sheet; the next product's two rows don't fit the rest
    write_items([single, SHOPIFY_PRODUCT, single, SHOPIFY_PRODUCT], path)
    sheets, book = read_sheets(path)
    assert list(sheets) == ["items", "items 2"]
    assert [r["id"] for r in sheets["items"]] == ["1", "7", "7"]
    assert [r["id"] for r in sheets["items 2"]] == ["1", "7", "7"]
    assert book["items 2"].freeze_panes == "A2"
    assert book["items 2"].auto_filter.ref == book["items"].auto_filter.ref


def test_item_longer_than_a_sheet_splits_across_sheets(tmp_path, monkeypatch):
    monkeypatch.setattr(xlsx_module, "MAX_ROWS", 2)  # header + 1 row
    path = tmp_path / "out.xlsx"
    write_items([SHOPIFY_PRODUCT], path)
    sheets, _ = read_sheets(path)
    assert [[r["variant_id"] for r in rows] for rows in sheets.values()] == [["71"], ["72"]]


def test_other_dicts_keep_their_types(tmp_path):
    path = tmp_path / "out.xlsx"
    rows = [
        {"title": "A", "min_price": Decimal("19.90"), "n": 3, "tags": ["a", "b"], "x": None},
        {"title": "B", "min_price": None, "n": 10**16, "tags": {"k": 1}, "x": float("nan")},
    ]
    write_items(rows, path)
    first, second = read_rows(path)
    assert first == {"title": "A", "min_price": 19.9, "n": 3, "tags": "a,b", "x": None}
    # Numbers Excel can't hold exactly are text
    assert (second["n"], second["x"], second["tags"]) == (str(10**16), "nan", '{"k": 1}')


def test_fields_pick_columns(tmp_path):
    path = tmp_path / "out.xlsx"
    write_items([SHOPIFY_PRODUCT], path, fields=["title", "price"])
    assert read_rows(path) == [{"title": "Hat", "price": 10}, {"title": "Hat", "price": 12.5}]


def test_no_items_no_file(tmp_path):
    path = tmp_path / "out.xlsx"
    assert write_items([], path) == 0
    assert not path.exists()


def test_append_rejected(tmp_path):
    with pytest.raises(ValueError, match="append"):
        ItemWriter(tmp_path / "x.xlsx", append=True)


def test_missing_xlsxwriter_fails_before_any_item(tmp_path, monkeypatch):
    monkeypatch.setitem(sys.modules, "xlsxwriter", None)
    with pytest.raises(ModuleNotFoundError, match=r"shopify-spy\[excel\]"):
        ItemWriter(tmp_path / "x.xlsx")


def test_lone_surrogates_dont_lose_the_file(tmp_path):
    # json.loads accepts a broken escape; xlsxwriter can't encode it
    path = tmp_path / "out.xlsx"
    bad = json.loads('{"title": "bad \\ud83d here", "n": 1}')
    long = {"title": "😀" + "\ud83d" + "x" * 40_000, "n": 2}
    write_items([{"title": "ok", "n": 0}, bad, long], path)
    rows = read_rows(path)
    assert [r["title"] for r in rows[:2]] == ["ok", "bad � here"]
    assert rows[2]["title"].startswith("😀�x")


def test_failed_close_removes_the_partial_file(tmp_path, monkeypatch):
    path = tmp_path / "out.xlsx"

    def fail(self):
        path.write_bytes(b"PK partial")
        raise OSError("disk full")

    monkeypatch.setattr(xlsx_module.XlsxWriter, "close", fail)
    with pytest.raises(OSError, match="disk full"):
        write_items([SHOPIFY_PRODUCT], path)
    assert not path.exists()


def test_summary_prices_are_numbers(tmp_path):
    path = tmp_path / "out.xlsx"
    write_items([summarize(WOO_PRODUCT).asdict()], path)
    (row,) = read_rows(path)
    assert (row["min_price"], row["max_price"]) == (19.99, 29.99)


def test_numbers_excel_would_round_stay_text(tmp_path):
    path = tmp_path / "out.xlsx"
    write_items([{"a": Decimal("12345678901234.567"), "price": "0.12345678901234567"}], path)
    (row,) = read_rows(path)
    assert row == {"a": "12345678901234.567", "price": "0.12345678901234567"}


def test_datetime_values_are_dates(tmp_path):
    path = tmp_path / "out.xlsx"
    seen = datetime(2024, 1, 15, 10, 30, tzinfo=timezone(timedelta(hours=-5)))
    write_items([{"seen": seen, "day": date(2024, 1, 2), "old": datetime(1800, 1, 1)}], path)
    (row,) = read_rows(path)
    assert row == {
        "seen": datetime(2024, 1, 15, 15, 30),
        "day": datetime(2024, 1, 2),
        "old": "1800-01-01 00:00:00",
    }


def test_out_of_range_values_stay_text(tmp_path):
    path = tmp_path / "out.xlsx"
    item = {
        "price": "1E+999999999",
        "compare_at_price": "1e-999999999",
        "created_at": "0001-01-01T00:00:00+05:00",
        "updated_at": "9999-12-31T23:00:00-05:00",
        "published_at": "1800-01-01T00:00:00-05:00",
    }
    write_items([item], path)
    (row,) = read_rows(path)
    assert row == item


def test_unwritable_path_fails_at_the_first_item(tmp_path):
    (tmp_path / "out.xlsx").mkdir()
    items_seen = []

    def items():
        for item in [SHOPIFY_PRODUCT, SHOPIFY_PRODUCT]:
            items_seen.append(item)
            yield item

    with pytest.raises(OSError):
        write_items(items(), tmp_path / "out.xlsx")
    assert len(items_seen) == 1


def test_temp_files_live_next_to_the_output_and_are_removed(tmp_path, monkeypatch):
    path = tmp_path / "out.xlsx"
    with ItemWriter(path) as writer:
        writer.write(SHOPIFY_PRODUCT)
        (tmp,) = [p for p in tmp_path.iterdir() if p.is_dir()]
        assert tmp.name.startswith("out.xlsx.tmp-")
    assert [p.name for p in tmp_path.iterdir()] == ["out.xlsx"]

    # Also when writing the workbook fails
    def fail(self):
        raise OSError("disk full")

    monkeypatch.setattr(xlsxwriter.Workbook, "close", fail)
    with pytest.raises(OSError, match="disk full"):
        write_items([SHOPIFY_PRODUCT], path)
    assert list(tmp_path.iterdir()) == []


def test_failed_temp_folder_leaves_no_empty_file(tmp_path, monkeypatch):
    def fail(*args, **kwargs):
        raise OSError("File name too long")

    monkeypatch.setattr(xlsx_module.tempfile, "TemporaryDirectory", fail)
    path = tmp_path / "out.xlsx"
    with pytest.raises(OSError, match="too long"):
        write_items([SHOPIFY_PRODUCT], path)
    assert list(tmp_path.iterdir()) == []


def test_subnormal_numbers_stay_text(tmp_path):
    path = tmp_path / "out.xlsx"
    write_items([{"price": "1e-310", "x": 1e-310, "compare_at_price": "0.000"}], path)
    (row,) = read_rows(path)
    assert row == {"price": "1e-310", "x": "1e-310", "compare_at_price": 0}


def test_failed_close_closes_the_temp_files(tmp_path, monkeypatch):
    # Windows can't delete the temp folder while xlsxwriter's row files are open
    writer = xlsx_module.XlsxWriter(tmp_path / "out.xlsx", ["a"])
    writer.write([["x"]])
    sheets = writer._book.worksheets()

    def fail(self):
        raise OSError("disk full")

    monkeypatch.setattr(xlsxwriter.Workbook, "close", fail)
    with pytest.raises(OSError):
        writer.close()
    assert all(sheet.row_data_fh_closed for sheet in sheets)


def test_typed_columns_are_unified_columns():
    # A renamed column would otherwise turn into text cells without notice
    assert ID_COLUMNS | NUMBER_COLUMNS | DATE_COLUMNS <= set(COLUMNS)
