"""End-to-end Excel output: the real CLI writes -F xlsx for local stores.

Run with: uv run pytest -m e2e
"""

import http.server
import json
import subprocess
import sys
import threading
from datetime import datetime

import openpyxl
import pytest

from shopify_spy._flatten import COLUMNS

from .test_e2e_diagnosis import _free_port

pytestmark = pytest.mark.e2e

SHOPIFY_PRODUCTS = [
    {
        "id": 8123456789012,  # 13 digits: Excel would show 8.12346E+12 as a number
        "title": "=HYPERLINK(\"http://evil\")",  # store text that looks like a formula
        "handle": "hat",
        "vendor": "Acme",
        "tags": ["wool"],
        "created_at": "2024-01-15T10:30:00-05:00",
        "body_html": "<p>Café \U0001f600</p>",
        "options": [{"name": "Size"}],
        "images": [],
        "variants": [
            {"id": 45123456789012, "title": "S", "option1": "S", "sku": "00123",
             "price": "10.00", "compare_at_price": None, "available": True, "grams": 200},
            {"id": 45123456789013, "title": "M", "option1": "M", "sku": "00124",
             "price": "12.50", "compare_at_price": "15.00", "available": False, "grams": 210},
        ],
    },
    {"id": 8123456789013, "title": "Gift card", "handle": "gift", "variants": []},
]  # fmt: skip

WOO_PRODUCTS = [
    {
        "id": 501,
        "name": "Pan",
        "slug": "pan",
        "permalink": "http://shop/product/pan/",
        "sku": "P1",
        "prices": {"price": "1990", "regular_price": "2490", "currency_code": "EUR",
                   "currency_minor_unit": 2, "price_range": None},
        "average_rating": "4.50",
        "review_count": 3,
        "is_in_stock": True,
        "images": [{"src": "http://shop/p.jpg"}],
        "attributes": [],
        "variations": [],
    }
]  # fmt: skip


class _Store(http.server.BaseHTTPRequestHandler):
    def log_message(self, *args):
        pass

    def do_GET(self):
        path, _, query = self.path.partition("?")
        first_page = "page=1" in query.split("&") or "page" not in query
        if path == "/products.json":
            body = {"products": SHOPIFY_PRODUCTS if first_page else []}
        elif path == "/wp-json/wc/store/v1/products":
            body = WOO_PRODUCTS if first_page else []
        else:
            self.send_response(404)
            self.send_header("Content-Length", "0")
            self.end_headers()
            return
        data = json.dumps(body).encode()
        self.send_response(200)
        self.send_header("Content-Type", "application/json")
        self.send_header("Content-Length", str(len(data)))
        self.end_headers()
        self.wfile.write(data)


@pytest.fixture
def store():
    server = http.server.ThreadingHTTPServer(("127.0.0.1", _free_port()), _Store)
    server.daemon_threads = True
    threading.Thread(target=server.serve_forever, daemon=True).start()
    yield f"http://127.0.0.1:{server.server_address[1]}"
    server.shutdown()


def _scrape_xlsx(tmp_path, url, *args):
    result = subprocess.run(
        [sys.executable, "-m", "shopify_spy.cli", "scrape", url, "-F", "xlsx",
         "--output", str(tmp_path), "--quiet", *args],
        capture_output=True, text=True, timeout=120, cwd=tmp_path,
    )  # fmt: skip
    assert result.returncode == 0, result.stdout + result.stderr
    (path,) = tmp_path.glob("data/*.xlsx")
    # No temp folder left next to the file
    assert list(path.parent.iterdir()) == [path]
    status = json.loads(next(tmp_path.glob("status/*_status.json")).read_text())
    assert status["output_error"] is None
    assert status["items_written"] == status["items_scraped"]
    book = openpyxl.load_workbook(path)
    sheet = book.active
    header, *rows = sheet.iter_rows(values_only=True)
    assert list(header) == [c for c in COLUMNS if c != "image_paths"]
    assert sheet.freeze_panes == "A2"
    assert (
        sheet.auto_filter.ref
        == f"A1:{sheet.cell(row=len(rows) + 1, column=len(header)).coordinate}"
    )
    return sheet, [dict(zip(header, row, strict=True)) for row in rows]  # fmt: skip


def test_shopify_store_to_xlsx(tmp_path, store):
    sheet, rows = _scrape_xlsx(tmp_path, store)

    # Two variant rows for the hat, one row for the gift card without variants
    assert len(rows) == 3
    hat_s, hat_m, gift = rows
    assert (hat_s["platform"], hat_s["kind"]) == ("shopify", "product")
    # IDs are text, so Excel doesn't show them in scientific notation
    assert (hat_s["id"], hat_s["variant_id"], gift["id"]) == (
        "8123456789012",
        "45123456789012",
        "8123456789013",
    )
    assert (hat_s["price"], hat_m["price"], hat_m["compare_at_price"]) == (10, 12.5, 15)
    assert (hat_s["compare_at_price"], hat_s["available"], hat_m["available"]) == (
        None,
        True,
        False,
    )
    assert (hat_s["grams"], hat_s["variant_count"], gift["variant_count"]) == (200, 2, 0)
    assert hat_s["sku"] == "00123"
    assert hat_s["created_at"] == datetime(2024, 1, 15, 15, 30)
    assert hat_s["description"] == "Café \U0001f600"
    # The formula-looking title is a string cell, not a formula
    title_col = list(next(sheet.iter_rows(max_row=1, values_only=True))).index("title") + 1
    cell = sheet.cell(row=2, column=title_col)
    assert (cell.value, cell.data_type) == ('=HYPERLINK("http://evil")', "s")
    for col in ("id", "variant_id", "price", "available", "created_at"):
        cell = sheet.cell(row=2, column=COLUMNS.index(col) + 1)
        assert cell.data_type == {"price": "n", "available": "b", "created_at": "d"}.get(
            col, "s"
        ), col


def test_woocommerce_store_to_xlsx(tmp_path, store):
    _, rows = _scrape_xlsx(tmp_path, store, "-p", "woocommerce")

    (row,) = rows
    assert (row["platform"], row["id"], row["title"], row["currency"]) == (
        "woocommerce",
        "501",
        "Pan",
        "EUR",
    )
    assert (row["price"], row["compare_at_price"], row["min_price"]) == (19.9, 24.9, 19.9)
    assert (row["average_rating"], row["review_count"], row["available"]) == (4.5, 3, True)
